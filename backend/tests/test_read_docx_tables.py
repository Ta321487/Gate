"""docx 抽取须含表格正文（任务书常把题目/要求写在表内）。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.services.proposal import _read_docx, business_signal_score
from app.services.upload_cluster import build_profiles, rule_cluster


class ReadDocxTablesTests(unittest.TestCase):
    def test_table_only_taskbook_not_empty(self):
        """正文几乎全在表格时，不得只抽出标题行。"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "方木-任务书.docx"
            doc = Document()
            doc.add_paragraph("毕业设计（论文）任务书")
            table = doc.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "设计题目:面向特定场景的健身管理平台的设计与实现"
            table.cell(1, 0).text = (
                "1.本课题的目的与意义\n"
                "本课题主要目的是设计并实现一个健身管理平台，"
                "前台用户使用 Vue.js 框架，后台使用 SpringBoot 与 MySQL，"
                "完成商品浏览、购物车下单、订单管理与后台信息管理。"
            )
            doc.save(str(path))

            text = _read_docx(path)
            self.assertIn("健身管理平台", text)
            self.assertIn("购物车", text)
            self.assertGreater(len(text), 80)
            self.assertGreaterEqual(business_signal_score(text), 5)

            profiles = build_profiles([(path, path.name, path.stat().st_size)])
            self.assertEqual(len(profiles), 1)
            self.assertEqual(profiles[0].role, "taskbook")
            clusters, discard, _notes = rule_cluster(profiles)
            self.assertEqual(discard, [])
            self.assertEqual(len(clusters), 1)


if __name__ == "__main__":
    unittest.main()
